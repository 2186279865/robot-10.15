from pathlib import Path
from copy import deepcopy
import json, hashlib, math, shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE = Path('C:/Users/rs/Desktop/附件7-命题文档_1785202024444.docx')
DEST = HERE.parent / '几何物块智能分拣系统作品说明书.docx'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == '7c49e14664f3795d0f7ec58d0b5f1f59784145c8e426350145f9a7dabfd03245'
FONT = 'C:/Windows/Fonts/msyh.ttc'

def font(n, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc' if bold else FONT, n)

def centered(draw, box, lines, size=25, bold=False):
    x0,y0,x1,y1=box
    lines=lines.split('\n')
    lineh=size*1.5
    y=(y0+y1-lineh*len(lines))/2
    for line in lines:
        draw.text(((x0+x1-draw.textlength(line,font=font(size,bold)))/2,y),line,font=font(size,bold),fill='#17202b')
        y+=lineh

def arrow(d,a,b):
    d.line([a,b],fill='#5d6976',width=4)
    ang=math.atan2(b[1]-a[1],b[0]-a[0]);l=15
    d.polygon([b,(b[0]-l*math.cos(ang-.45),b[1]-l*math.sin(ang-.45)),(b[0]-l*math.cos(ang+.45),b[1]-l*math.sin(ang+.45))],fill='#5d6976')

def box(d,coords,text,size=25):
    d.rounded_rectangle(coords,radius=12,fill='#eef3f7',outline='#5d6976',width=2)
    centered(d,coords,text,size)

def architecture():
    im=Image.new('RGB',(1500,620),'white');d=ImageDraw.Draw(im)
    box(d,(30,45,340,165),'双路相机\n顶视定位  副视补充')
    box(d,(30,280,340,410),'香橙派\n图像采集与串口转发')
    box(d,(500,45,1000,200),'笔记本上位机\n分割  分类  标定  任务调度')
    box(d,(1130,45,1470,200),'视觉模型\n云端主路径  本地备选')
    box(d,(500,285,1000,415),'STM32F407\n协议处理与闭环运动')
    box(d,(1130,285,1470,415),'EMM V5 与舵机\nXYZ  转盘  夹爪')
    box(d,(500,490,1000,590),'显示终端\n类别  盒号  件数  状态')
    arrow(d,(185,165),(185,280));arrow(d,(340,325),(500,125));arrow(d,(1000,125),(1130,125));arrow(d,(1130,160),(1000,160));arrow(d,(340,365),(500,350));arrow(d,(750,200),(750,285));arrow(d,(1000,350),(1130,350));arrow(d,(750,415),(750,490))
    im.save(HERE/'architecture.png')

def visual_flow():
    im=Image.new('RGB',(1500,650),'white');d=ImageDraw.Draw(im)
    labels=['双摄原图\n固定机位与工作区','四点透视校正\n统一托盘坐标','前景候选与 SAM\n轮廓  中心  面积','双摄位置颜色匹配\n补充侧面证据','视觉模型判断\n形状与主体颜色','按中心匹配标签\n失败走几何兜底']
    coords=[(35,40,460,185),(535,40,960,185),(1035,40,1460,185),(1035,285,1460,430),(535,285,960,430),(35,285,460,430)]
    for c,t in zip(coords,labels):box(d,c,t,27)
    for a,b in [((460,112),(535,112)),((960,112),(1035,112)),((1248,185),(1248,285)),((1035,357),(960,357)),((535,357),(460,357))]:arrow(d,a,b)
    box(d,(260,520,1240,620),'输出目标位置与语义结果  →  界面显示与分拣任务',28)
    arrow(d,(247,430),(247,570));arrow(d,(247,570),(260,570))
    im.save(HERE/'visual-flow.png')

def motion_flow():
    im=Image.new('RGB',(1500,500),'white');d=ImageDraw.Draw(im)
    texts=['识别与定位\n确定类别和盒号','转盘对准\n预备夹爪','低位移出通道\n抬升至安全高度','XY 对准目标\n下探并夹取','抬升与复核\n原路返回','P0 松爪投放\n显示后处理下一件']
    for i,t in enumerate(texts):
        col=i if i<3 else 5-i;row=0 if i<3 else 1
        x=30+col*500;y=35+row*260
        box(d,(x,y,x+440,y+160),t,28)
    for a,b in [((470,115),(530,115)),((970,115),(1030,115)),((1250,195),(1250,295)),((1030,375),(970,375)),((530,375),(470,375))]:arrow(d,a,b)
    im.save(HERE/'motion-flow.png')

architecture();visual_flow();motion_flow()
raw=Image.open(ROOT/'04-队友视觉调试包/samples/real_new1.jpg').convert('RGB')
w,h=raw.size
# 只呈现工作区，避开画面外围的人员、屏幕及标识。
raw.crop((int(w*.23),int(h*.15),int(w*.79),int(h*.88))).save(HERE/'workspace.jpg',quality=94)

shutil.copyfile(SOURCE,DEST)
doc=Document(DEST)
main=doc.tables[0]
# 保留主表；智能搬运专用空白场地图不属于本作品。
for element in list(doc._element.body):
    if element is main._tbl or element.tag==qn('w:sectPr'):continue
    doc._element.body.remove(element)
for style_name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    if style_name not in doc.styles:
        doc.styles.add_style(style_name,WD_STYLE_TYPE.PARAGRAPH)
    s=doc.styles[style_name]
    s.font.name='宋体';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    s.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    s.paragraph_format.line_spacing=Pt(20)
    s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(0)
    for border in list(s.element.xpath('.//w:pBdr')):border.getparent().remove(border)

def fmt(p,*,bold=False,indent=True,center=False,keep=False,size=10.5):
    f=p.paragraph_format;f.line_spacing_rule=WD_LINE_SPACING.EXACTLY;f.line_spacing=Pt(20)
    f.space_before=Pt(0);f.space_after=Pt(0);f.first_line_indent=Pt(21 if indent else 0)
    f.keep_with_next=keep;f.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name='宋体';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
        rf=r._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'宋体')
    for e in list(p._p.xpath('./w:pPr/w:pBdr')):e.getparent().remove(e)

def clear(cell):
    for x in list(cell._tc):
        if x.tag!=qn('w:tcPr'):cell._tc.remove(x)
    cell.add_paragraph()
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

stats={}
current=[]
def p(cell,text,bold=False,indent=True):
    para=cell.paragraphs[0] if len(cell.paragraphs)==1 and not cell.paragraphs[0].text else cell.add_paragraph()
    para.add_run(text);fmt(para,bold=bold,indent=indent,keep=bold)
    current.append(text)
    return para

def sub(cell,title):return p(cell,title,bold=True,indent=False)
def picture(cell,name,caption,width=13):
    para=cell.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
    para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.keep_with_next=True
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(2)
    para.add_run().add_picture(str(HERE/name),width=Cm(width))
    q=cell.add_paragraph(caption);fmt(q,center=True,indent=False,size=10.5)

def start(row,title,newpage=False):
    global current
    current=[]
    c=main.cell(row,0);clear(c)
    for e in list(main.rows[row]._tr.xpath('./w:trPr/w:trHeight')):e.getparent().remove(e)
    q=p(c,title,bold=True,indent=False)
    q.paragraph_format.page_break_before=False
    return c

def done(section,limit):
    count=sum(len(''.join(x.split())) for x in current)
    stats[section]={'nonspace_characters':count,'limit':limit}
    assert count<=limit,(section,count,limit)

# 保留赛事编号空白与源赛事标题，填入参赛项目。
main.cell(3,1).text='“智能＋”赛道\n智能分拣'
for q in main.cell(3,1).paragraphs:fmt(q,indent=False,center=True)
title=main.cell(2,0).paragraphs[0];title.style=doc.styles['Title'];fmt(title,bold=True,indent=False,center=True,size=22)

c=start(4,'一、参赛作品简介')
q=p(c,'几何物块视觉识别与自动分拣系统',bold=True,indent=False);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p(c,'本作品面向多形状、多颜色物块的智能分拣任务，采用双摄像头、笔记本、香橙派与STM32F407组成感知及控制系统，以X、Y、Z直线运动机构、六盒转盘和舵机夹爪完成物块搬运。系统围绕识别、定位、抓取、归盒与信息显示构建作业流程，工作区样例如图1所示。')
picture(c,'workspace.jpg','图1 项目相机采集的物块与抓取工作区',width=5.1)
p(c,'视觉端利用SAM提取物块轮廓及位置，结合通用视觉模型判断几何形状与颜色，无需为每批物块重新训练专用分类器；几何规则作为辅助。上位机将图像坐标转换为机构目标坐标，按类别分配储物盒，并调度转盘、夹爪及三轴动作。')
p(c,'系统提供视觉预览、手动控制、标定示教、自动抓放和日志记录。现阶段已具备软件与控制链路基础，正在完善复杂几何体及镂空件识别、全场标定和整机连续作业验证。总体架构如图2所示。')
picture(c,'architecture.png','图2 系统总体架构示意图',width=11.5)
done('一',500)

c=start(5,'二、作品设计方案',True)
p(c,'作品采用分层设计，将图像理解、任务决策与底层运动分开。笔记本承担分割、分类、坐标换算及任务管理；香橙派承担双摄采集和通信转发；STM32负责解析指令、访问驱动器并执行运动。六盒转盘将类别选择集中在固定投放位置，减少末端在多个投放点之间的移动。')
sub(c,'2.1 硬件设计方案')
sub(c,'2.1.1 双摄感知与计算模块')
p(c,'主视相机覆盖抓取区，提供物块数量、轮廓与平面位置；副视相机补充侧面信息，辅助区分投影相近的几何体。相机可直接接入笔记本，也可由香橙派通过图传链路提供图像。识别在笔记本运行，避免将较重的视觉推理压在采集端。')
sub(c,'2.1.2 控制与驱动模块')
p(c,'STM32F407通过USB虚拟串口与上位机通信，再经USART向EMM V5驱动器发送运动及查询指令。执行端包含Z升降轴、X和Y平移轴、转盘轴以及舵机夹爪。上位机利用位置回读检查动作完成情况，舵机通过PWM控制预备、夹紧与松开。实际回读接线、轴向和行程须在整机验收时逐项确认。')
sub(c,'2.1.3 转盘与末端执行机构')
p(c,'转盘设置六个储物盒，各盒相隔60°，物块由夹爪在固定位置P0投放。X、Y轴负责对准，Z轴负责接近和提升。初始位置附近存在通道约束，出程先沿Y轴离开盒区再抬升，抓取区内的横向移动在已验证的安全高度进行；回程按相反顺序返回。')
sub(c,'2.2 软件与视觉算法')
sub(c,'2.2.1 托盘校正与实例分割')
p(c,'两路图像分别选四点建立透视变换，将工作区校正到统一托盘平面。当前交付版先以颜色前景生成候选区域，再向SAM提供内部提示点，得到实例掩膜。面积、边缘和前景覆盖率过滤用于去除托盘及盘外干扰，并对重复掩膜或同一物块的相邻面进行合并。流程见图3。')
picture(c,'visual-flow.png','图3 交付版视觉识别流程示意图',width=12.5)
sub(c,'2.2.2 形状颜色识别与双摄协同')
p(c,'当前交付版配置以云端视觉模型为主分类器，输入校正后的主视图及可用副视图，输出物块形状、主体颜色和中心估计；再按中心邻近关系与SAM目标匹配，保留分割位置。云端失败时回退到轮廓圆度、锐角数、长宽比等几何判据及既有登记规则。双摄通过位置和颜色进行匹配，并非三维重建，同色重叠目标仍需复核。')
p(c,'新上位机另设本地通用视觉模型与云端裁剪分类入口，以掩膜颜色统计提供主体色参考，并增加可见孔洞分析。该路径用于评估首次出现物块的识别能力，不将单幅轮廓直接等同于完整三维形状。镂空件还需区分真实孔洞、印字和阴影，未确定结果不得作为可靠分类结论。')
sub(c,'2.2.3 坐标标定与目标管理')
p(c,'视觉中心使用托盘归一化坐标u、v表示，通过多点仿射标定转换为X、Y目标脉冲：')
q=p(c,'X＝a·u＋b·v＋tx    Y＝c·u＋d·v＋ty',indent=False);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p(c,'采样需覆盖工作区并使用新鲜的位置回读；拟合后以独立点验证误差，检查目标和路线是否处于各轴限位内。改变相机、工作区或机械零点后重新复核。现有标定尚存在局部误差及限位冲突，不能直接作为全场精度结论。物块高度引起的视差及镂空体抓点选择是后续完善重点。')
sub(c,'2.3 上下位机协同与自动作业')
p(c,'系统支持手动调试和自动抓放。自动流程先识别类别与位置，再确定盒号；转盘对准后，夹爪经通道到达安全高度，X、Y对准目标，Z下探夹取，抬升后返回P0松爪，最后更新类别、盒号和件数。新上位机按形状、颜色及镂空状态形成类别键，最多分配六盒；分类标准可按实际任务调整。动作顺序见图4。')
picture(c,'motion-flow.png','图4 自动抓放动作顺序示意图',width=12.5)
sub(c,'2.4 状态反馈与异常处理')
p(c,'通信包含请求编号、校验、应答和遥测，闭环定位由固件回读位置并补差。新上位机加入显式解除停止、过期位置拒绝、路线预检及完成后位置复核。旧流程的自动解锁、角度语义及部分时序问题已列入整改，需统一验收；原点、硬限位、独立故障保护和落料确认仍需结合实机完善。比赛通信及模型部署方式以现场规则为准。')
done('二',2000)

c=start(6,'三、制作与测试记录',True)
p(c,'项目按机构与驱动联调、图传与协议对接、视觉标定、自动流程集成推进。已形成STM32固件、香橙派服务和一体化上位机，并另建支持本地与云端识别的新界面。')
sub(c,'3.1 调试问题与处理')
p(c,'针对方块整体投影呈六边形的问题，加入顶面分析与视觉模型判断；针对托盘及重复面片误检，增加前景提示、边缘过滤和掩膜合并。标定检查发现残差偏大及目标越限，拟通过分散采点、独立验证和统一限位整改。连续分盒角度和停止取消问题已定位，新上位机采用绝对盒位及锁存取消。')
sub(c,'3.2 已有验证与后续测试')
p(c,'已归档的离线测试中，旧自动抓放、客户端及预备定位共56项单元测试通过。对一张含9件物块的渲染图，本地视觉模型形状判断为7件正确，热启动单件约1.2至1.6秒；该结果不代表现场准确率，且全图分割另有耗时。当前仍需统计实拍形状颜色准确率、标定毫米误差、抓放成功率及连续运行耗时，镂空件单独验收。')
done('三',500)

c=start(7,'四、创新点')
p(c,'1）分割与语义识别分工。由SAM提供物块边界与位置，通用视觉模型判断形状，面向无需逐批训练的应用需求，并保留几何辅助路径。',indent=False)
p(c,'2）固定投放点与六盒转盘协同。通过类别到盒位的映射，在统一投放位置完成分类收纳，结合低位通道和高位移动组织抓放路线。',indent=False)
p(c,'3）识别结果与执行状态联动。将目标时效、位置反馈、类别映射和显示记录纳入上位机；新实现以到位应答与位置复核约束后续动作，便于追溯异常。',indent=False)
done('四',300)

c=start(8,'五、决赛命题设想',True)
p(c,'建议以未知几何物块的自主识别与分类收纳为决赛场景，考查首次识别、多目标处理、定位抓取和状态反馈。以下物料数量、时间及评分为命题建议，正式任务与通信限制以组委会公布为准。')
sub(c,'5.1 决赛任务与阶段区别')
sub(c,'5.1.1 第一阶段基础任务')
p(c,'设置6件不超过3类的常见实心物块，分开放置，颜色和形状组合在准备阶段公布。装置完成识别、定位、抓取与归盒，显示序号、形状、颜色、盒号和完成状态。建议限时4分钟，重点评价基础分类、抓取成功及完整作业流程。')
sub(c,'5.1.2 第二阶段综合任务')
p(c,'设置12件、不超过6类的物块，分2、4、6件三批放入；增加此前未见的颜色与形状组合、不同转角、相邻摆放及至少2件镂空物块。准备结束后不允许重新训练模型或逐件人工登记。建议总限时7分钟，重点评价泛化、连续作业及对不确定目标的处置。')
p(c,'第一阶段采用分散、已知组合验证基本能力；第二阶段增加未知组合、批量目标及孔洞特征，同时要求批次统计与异常记录。类别判据应提前明确：可按形状加颜色分盒，也可增加镂空状态，但同一轮形成的类别键不得超过盒数。')
sub(c,'5.1.3 判定与约束')
p(c,'物块形状、颜色和镂空状态由裁判建立真值，尺寸、质量、可夹持面及孔洞可见性在准备阶段确认，避免将不可观测内部结构作为唯一判据。同一类别应进入同一盒，不同类别不得混放。识别显示、实际落盒和批次计数分别核验；执行过松爪动作不能单独算作落盒成功。')
p(c,'计时期间由装置自主决策，不允许人工指定类别和运动目标。允许的初始化与调整仅在准备阶段完成。若禁用外网，应使用预先部署的本地模型；若允许云端，应统一网络及调用条件。安全停止、人工介入和恢复的计分规则应在赛前明确。')
sub(c,'5.2 命题变化后的设计调整')
sub(c,'5.2.1 视觉识别调整')
p(c,'对首次出现的形状组合，以通用模型及双视角证据辅助判断；彩色目标与低饱和目标分别配置前景策略，避免黑白灰物块在候选阶段遗漏。分类低可信或两视角冲突时重新采图，并记录原因，不能直接映射到某个已知类。镂空检测使用孔洞形态与另一视角复核，抓点结合实际夹持面选择。')
sub(c,'5.2.2 标定与控制调整')
p(c,'按有效抓取区重新做九点或十六点标定，另留独立检验点。根据实际物块高度评估视差和抓取高度，必要时做分层补偿。统一绝对目标限位、盒号角度和运动完成条件；停止后取消旧任务，不自动恢复。原点、终端保护和夹持检测的补充需结合机构条件实施。')
sub(c,'5.2.3 效率与部署调整')
p(c,'将快速定位和语义分类异步执行，对稳定目标缓存结果，减少重复推理。SAM全图分割用于需要时的补充，不能把相机预览帧率当作分类速度。通过同一批实拍样本比较本地与云端的正确率、完整耗时和失败率，再选择比赛部署路径；优先保持判定一致与动作可靠。')
sub(c,'5.3 现场决赛过程')
p(c,'准备阶段建议5分钟：核对机械固定、供电和通信，清空储物盒，确认各轴位置基准与转盘1号盒零点，以固定参考点复核标定，检查停止功能，载入任务配置后进入待机。上述检查完成后才允许计时启动。')
p(c,'比赛开始后，裁判按批次放入物块并撤离工作区。系统采图、识别并建立类别映射，选择可抓目标，完成转盘对准、通道移动、抓取、返回和归盒。每件完成后更新显示，批次完成后显示类别及数量，裁判确认后投入下一批。遇到遮挡或冲突可重新观察；连续失败或越限时停止并保留故障记录。')
p(c,'结束阶段由裁判清点各盒实物，与识别记录和完成计数对照，统计错分、漏分、抓取失败及人工介入。盒区当前没有独立视觉确认条件，因此不将自主回收纠错列为已具备能力。若未来新增落料传感器或盒区相机，可另设复核任务。')
sub(c,'5.4 建议评分与排序')
p(c,'建议分类正确性40分，按形状、颜色、镂空及实际归盒一致性计分；自主完成度30分，按有效搬运数量和完整批次计分；时间效率20分，仅对有效完成任务计分；交互与异常处理10分，评价显示、记录和受控停止。总分相同时，先比较正确归盒数量，再比较有效完成时间。具体细则由赛题统一规定。')
done('五',2000)

# 每个内容单元自然排版，取消模板占位行高和不必要的整行不跨页设置。
for i in range(4,10):
    row=main.rows[i]
    for e in list(row._tr.xpath('./w:trPr/w:cantSplit')):e.getparent().remove(e)
    allow=OxmlElement('w:cantSplit');allow.set(qn('w:val'),'0');row._tr.get_or_add_trPr().append(allow)
for i in range(4,9):
    cell=main.cell(i,0)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    # 小幅增添段落上下留白，正文行距仍严格20磅。
    for para in cell.paragraphs:
        if para.text:
            para.paragraph_format.keep_together=True
        if para.text and para.runs and para.runs[0].bold and para is not cell.paragraphs[0]:
            para.paragraph_format.space_before=Pt(2)
for i,q in enumerate(main.cell(9,0).paragraphs):
    fmt(q,indent=False,size=10.5,keep=i<len(main.cell(9,0).paragraphs)-1)
    q.paragraph_format.keep_together=True
main.rows[9]._tr.get_or_add_trPr().find(qn('w:cantSplit')).set(qn('w:val'),'1')

# Word对长合并行的分页不稳定；按标题正文或图片图题成组克隆同一模板行。
# 组间隐藏横线，章节外框及章间分隔保持源表视觉样式。
for row in list(main.rows)[4:9]:
    tc=row._tr.tc_lst[0]
    paras=list(tc.xpath('./w:p'))
    groups=[]; group=[]
    for para in paras:
        group.append(para)
        keep=para.xpath('./w:pPr/w:keepNext')
        if not keep or keep[0].get(qn('w:val')) in ('0','false'):
            groups.append(group);group=[]
    if group:groups.append(group)
    for gi,group in enumerate(groups):
        tr=deepcopy(row._tr)
        for e in list(tr.xpath('./w:tblPrEx')):tr.remove(e)
        trpr=tr.get_or_add_trPr()
        for e in list(trpr):trpr.remove(e)
        trpr.append(OxmlElement('w:cantSplit'))
        newtc=tr.tc_lst[0]
        for e in list(newtc):
            if e.tag!=qn('w:tcPr'):newtc.remove(e)
        tcpr=newtc.get_or_add_tcPr()
        for e in list(tcpr.xpath('./w:tcBorders')):tcpr.remove(e)
        borders=OxmlElement('w:tcBorders')
        for side,visible in [('top',gi==0),('bottom',gi==len(groups)-1)]:
            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single' if visible else 'nil')
            if visible:edge.set(qn('w:sz'),'8');edge.set(qn('w:color'),'auto')
            borders.append(edge)
        tcpr.append(borders)
        for para in group:
            clone=deepcopy(para)
            for keep in clone.xpath('./w:pPr/w:keepNext'):keep.set(qn('w:val'),'0')
            newtc.append(clone)
        row._tr.addprevious(tr)
    main._tbl.remove(row._tr)
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.title='几何物块视觉识别与自动分拣系统作品说明书'
doc.core_properties.subject='智能分拣作品说明书'
doc.core_properties.keywords='';doc.core_properties.comments=''
doc.save(DEST)
(HERE/'word-counts.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == '7c49e14664f3795d0f7ec58d0b5f1f59784145c8e426350145f9a7dabfd03245'
print(DEST)
print(json.dumps(stats,ensure_ascii=False))
